"""Caixa de entrada: recebe documentos (upload hoje; Drive/WhatsApp/e-mail depois), identifica o tipo
(boleto, nota fiscal, relatório de vendas) e encaminha para a ação certa.

- NF-e em XML: lida por inteiro (emitente, itens, parcelas) e vira compra + contas a pagar.
- Boleto e NF em PDF/foto: abertos no leitor de boleto (boleto.processar) para conferência.
- Relatórios de vendas (Anota Aí, Sischef, Aiqfome): guardados com prévia; a leitura vem na próxima etapa.
"""
import csv
import datetime as dt
import hashlib
import io
import logging
import os
import re
import uuid
import zipfile
from decimal import Decimal

import pdfplumber
import psycopg
from defusedxml import ElementTree as ET
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from pydantic import BaseModel, Field

import boleto
import fin
from fin import DSN, _conv, q

log = logging.getLogger("inbox")
router = APIRouter(prefix="/api/inbox")
DOCS_DIR = os.path.join(boleto.ANEXOS_DIR, "docs")
MAX_BYTES = 20 * 1024 * 1024
NS = {"n": "http://www.portalfiscal.inf.br/nfe"}
FONTES = {"anota_ai": ("anota ai", "anota aí", "anotaai"), "sischef": ("sischef",), "aiqfome": ("aiqfome",)}
ROT_FONTE = {"anota_ai": "Anota Aí", "sischef": "Sischef", "aiqfome": "Aiqfome"}


def _brl(v) -> str:
    return "R$ " + f"{float(v or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _dm(d) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


# ------------------------------------------------------------------ identificação
def _tipo_arquivo(conteudo: bytes, nome: str) -> tuple[str, str]:
    """(mime, extensão)"""
    cab = conteudo[:200].lstrip()
    n = nome.lower()
    if conteudo[:5] == b"%PDF-":
        return "application/pdf", "pdf"
    if cab.startswith(b"<?xml") or b"<nfeProc" in conteudo[:2000] or b"<NFe" in conteudo[:2000] or n.endswith(".xml"):
        return "application/xml", "xml"
    if conteudo[:2] == b"PK" and (n.endswith(".xlsx") or b"xl/" in conteudo[:4000]):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    if n.endswith(".csv") or n.endswith(".txt"):
        return "text/csv", "csv"
    try:
        mime = boleto._mime(conteudo, nome)
        return mime, {"image/jpeg": "jpg", "image/png": "png", "image/heic": "heic", "image/webp": "webp"}.get(mime, "img")
    except HTTPException:
        return "application/octet-stream", "bin"


def _fonte_vendas(texto: str):
    t = re.sub(r"[-_./]+", " ", texto.lower())   # "relatorio-anota-ai.csv" → "relatorio anota ai csv"
    return next((k for k, palavras in FONTES.items() if any(p in t for p in palavras)), None)


def _num(x):
    return Decimal(x) if x not in (None, "") else None


def ler_nfe(conteudo: bytes) -> dict | None:
    """Lê o XML da NF-e (modelo 55/65). Devolve None se não for uma NF-e."""
    try:
        raiz = ET.fromstring(conteudo)
    except Exception:
        return None
    inf = raiz.find(".//n:infNFe", NS)
    if inf is None:
        return None
    g = lambda el, path: (el.findtext(path, namespaces=NS) or "").strip() if el is not None else ""
    chave = re.sub(r"\D", "", inf.get("Id", "")) or re.sub(r"\D", "", g(raiz, ".//n:protNFe/n:infProt/n:chNFe"))
    ide, emit, tot = inf.find("n:ide", NS), inf.find("n:emit", NS), inf.find("n:total/n:ICMSTot", NS)
    emissao = (g(ide, "n:dhEmi") or g(ide, "n:dEmi"))[:10]
    itens = []
    for det in inf.findall("n:det", NS):
        p = det.find("n:prod", NS)
        itens.append({"n": int(det.get("nItem") or len(itens) + 1), "codigo": g(p, "n:cProd"),
                      "ean": g(p, "n:cEAN") if g(p, "n:cEAN") not in ("SEM GTIN", "") else None,
                      "descricao": g(p, "n:xProd"), "ncm": g(p, "n:NCM"), "cfop": g(p, "n:CFOP"),
                      "unidade": g(p, "n:uCom"), "quantidade": str(_num(g(p, "n:qCom")) or 0),
                      "valor_unitario": str(_num(g(p, "n:vUnCom")) or 0), "valor_total": str(_num(g(p, "n:vProd")) or 0)})
    parcelas = [{"numero": g(d, "n:nDup"), "vencimento": g(d, "n:dVenc")[:10], "valor": str(_num(g(d, "n:vDup")) or 0)}
                for d in inf.findall("n:cobr/n:dup", NS)]
    return {"chave": chave or None, "numero": g(ide, "n:nNF"), "serie": g(ide, "n:serie"),
            "emissao": emissao or None, "natureza": g(ide, "n:natOp"),
            "finalidade": g(ide, "n:finNFe") or "1", "tipo_debito": g(ide, "n:tpNFDebito") or None,
            "emitente": {"cnpj": g(emit, "n:CNPJ") or g(emit, "n:CPF"), "nome": g(emit, "n:xNome"),
                         "fantasia": g(emit, "n:xFant")},
            "valor_produtos": str(_num(g(tot, "n:vProd")) or 0), "valor_frete": str(_num(g(tot, "n:vFrete")) or 0),
            "valor_desconto": str(_num(g(tot, "n:vDesc")) or 0), "valor_total": str(_num(g(tot, "n:vNF")) or 0),
            "itens": itens, "parcelas": parcelas}


def _previa_planilha(conteudo: bytes, ext: str) -> dict:
    linhas = []
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 200:
                break
            linhas.append(["" if c is None else str(c) for c in row])
        total = ws.max_row
    else:
        texto = next((conteudo.decode(enc) for enc in ("utf-8-sig", "latin-1") if _decodifica(conteudo, enc)), "")
        try:
            dialeto = csv.Sniffer().sniff(texto[:4000], delimiters=";,\t|")
        except csv.Error:
            dialeto = csv.excel
        todas = list(csv.reader(io.StringIO(texto), dialeto))
        linhas, total = todas[:200], len(todas)
    linhas = [l for l in linhas if any(c.strip() for c in l)]
    return {"colunas": linhas[0] if linhas else [], "linhas": linhas[1:16], "total_linhas": max(total - 1, 0)}


def _decodifica(b, enc):
    try:
        b.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def classificar(conteudo: bytes, nome: str, mime: str, ext: str, pasta: str) -> dict:
    """Devolve tipo, subtipo, resumo e dados extraídos (e grava a prévia da 1ª página quando houver)."""
    if ext == "xml":
        nfe = ler_nfe(conteudo)
        if nfe:
            resumo = (f"NF {nfe['numero']} · {nfe['emitente']['nome']} · {_brl(nfe['valor_total'])} · "
                      f"{len(nfe['itens'])} itens · {len(nfe['parcelas']) or 1} parcela(s)")
            return {"tipo": "nota_fiscal", "subtipo": "nfe_xml", "resumo": resumo, "dados": {"nfe": nfe}}
        return {"tipo": "outro", "subtipo": "xml", "resumo": "XML que não é NF-e", "dados": {}}
    if ext in ("csv", "xlsx"):
        prev = _previa_planilha(conteudo, ext)
        fonte = _fonte_vendas(nome + " " + " ".join(prev["colunas"]) + " " + " ".join(" ".join(l) for l in prev["linhas"][:5]))
        rot = ROT_FONTE.get(fonte, "origem não identificada")
        return {"tipo": "relatorio_vendas", "subtipo": fonte, "dados": {"previa": prev},
                "resumo": f"Relatório de vendas · {rot} · {prev['total_linhas']} linhas"}
    if mime == "application/pdf" or mime.startswith("image/"):
        texto = ""
        if mime == "application/pdf":
            with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
                texto = " ".join((pg.extract_text() or "") for pg in pdf.pages[:3])
        paginas = None
        if len(texto.strip()) < 40:  # escaneado ou foto: OCR
            paginas = boleto._paginas(conteudo, mime)
            texto = " ".join(w["t"] for _, ws, _ in paginas for w in ws)
        # prévia da 1ª página
        try:
            img = (paginas[0][0] if paginas else boleto._paginas(conteudo, mime)[0][0]).copy()
            img.thumbnail((900, 1300))
            img.save(os.path.join(pasta, "previa.png"), optimize=True)
        except Exception:
            log.warning("sem prévia para %s", nome)
        linha = boleto.linha_digitavel(texto)
        tl = texto.lower()
        chave = re.search(r"(?<!\d)(\d{4}\s?){11}(?!\d)", texto)
        eh_nf = "danfe" in tl or "nota fiscal eletr" in tl or "nfc-e" in tl or "chave de acesso" in tl
        if eh_nf:
            m = re.search(r"valor\s+total\s+da\s+nota[^\d]{0,40}([\d.]+,\d{2})", texto, re.I)
            dados = {"chave": re.sub(r"\D", "", chave.group(0)) if chave else None,
                     "valor_total": m.group(1) if m else None, "linha": linha}
            return {"tipo": "nota_fiscal", "subtipo": "nfe_pdf", "dados": dados,
                    "resumo": "Nota fiscal (PDF/foto)" + (f" · {_brl(m.group(1).replace('.', '').replace(',', '.'))}" if m else "")
                              + (" · com boleto" if linha else "")}
        if linha:
            venc = dt.date.fromisoformat(linha["vencimento"]) if linha.get("vencimento") else None
            return {"tipo": "boleto", "subtipo": linha["tipo"], "dados": {"linha": linha},
                    "resumo": f"Boleto · {_brl(linha['valor'])} · vence {_dm(venc)}" if linha.get("valor") else "Boleto"}
        fonte = _fonte_vendas(nome + " " + texto[:3000])
        if fonte and re.search(r"venda|pedido|relat", tl):
            return {"tipo": "relatorio_vendas", "subtipo": fonte, "dados": {},
                    "resumo": f"Relatório de vendas · {ROT_FONTE[fonte]} (PDF)"}
        return {"tipo": "outro", "subtipo": None, "dados": {"texto": texto[:500]},
                "resumo": "Documento não identificado"}
    return {"tipo": "outro", "subtipo": None, "dados": {}, "resumo": "Formato não reconhecido"}


def receber(conteudo: bytes, nome: str, origem: str = "upload", origem_ref: str | None = None) -> dict:
    """Entrada única para todos os canais (upload, Drive, WhatsApp, e-mail)."""
    if not conteudo:
        raise HTTPException(422, f"{nome}: arquivo vazio")
    if len(conteudo) > MAX_BYTES:
        raise HTTPException(413, f"{nome}: arquivo maior que 20 MB")
    sha = hashlib.sha256(conteudo).hexdigest()
    existente = q("select id, nome, tipo, resumo, status from fin_app.documentos where sha256 = %s", (sha,))
    if existente:
        return {**existente[0], "duplicado": True}
    did = uuid.uuid4()
    pasta = os.path.join(DOCS_DIR, str(did))
    os.makedirs(pasta, exist_ok=True)
    mime, ext = _tipo_arquivo(conteudo, nome)
    with open(os.path.join(pasta, f"original.{ext}"), "wb") as f:
        f.write(conteudo)
    try:
        c = classificar(conteudo, nome, mime, ext, pasta)
    except Exception as e:
        log.exception("falha classificando %s", nome)
        c = {"tipo": "outro", "subtipo": None, "dados": {"erro": type(e).__name__}, "resumo": "Não consegui ler o arquivo"}
    with psycopg.connect(**DSN, row_factory=dict_row) as conn:
        row = conn.execute("""insert into fin_app.documentos (id, origem, origem_ref, nome, mime, tamanho, sha256, tipo,
                                  subtipo, resumo, dados)
                              values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                              returning id, nome, tipo, resumo, status""",
                           (did, origem, origem_ref, (nome or "documento")[:200], mime, len(conteudo), sha,
                            c["tipo"], c["subtipo"], c["resumo"], Jsonb(c["dados"]))).fetchone()
    return {**{k: _conv(v) for k, v in row.items()}, "duplicado": False}


# ------------------------------------------------------------------ endpoints
@router.post("")
def enviar(arquivos: list[UploadFile] = File(...)):
    out = []
    for a in arquivos[:30]:
        try:
            out.append(receber(a.file.read(MAX_BYTES + 1), a.filename or "documento"))
        except HTTPException as e:
            out.append({"nome": a.filename, "erro": e.detail})
    return out


@router.get("")
def listar(status: str = "novo", tipo: str = "", antigos: bool = False):
    # Processados: por padrão só os dos últimos 30 dias (nada é apagado; os antigos aparecem com antigos=true)
    recentes = status == "processado" and not antigos
    filtro = {"status": status, "tipo": tipo, "recentes": recentes}
    docs = q("""select id, origem, nome, mime, tamanho, tipo, subtipo, status, resumo, lancamento_ids,
                       recebido_em, processado_em,
                       exists (select 1 from fin_app.compras c where c.documento_id = d.id) as compra_lancada,
                       (select c.conciliada_em from fin_app.compras c where c.documento_id = d.id limit 1) as conciliada_em,
                       coalesce((select c.finalidade from fin_app.compras c where c.documento_id = d.id limit 1),
                                (select x.rot from (values ('2','Nota complementar'),('3','Nota de ajuste'),('4','Nota de devolução'),
                                   ('5','Nota de crédito'),('6','Nota de débito')) x(f, rot) where x.f = d.dados->'nfe'->>'finalidade')) as finalidade
                from fin_app.documentos d
                where (%(status)s = 'todos' or status = %(status)s) and (%(tipo)s = '' or tipo = %(tipo)s)
                  and (not %(recentes)s or coalesce(processado_em, recebido_em) >= now() - interval '30 days')
                order by recebido_em desc limit 300""", filtro)
    ocultos = q("""select count(*) as n from fin_app.documentos where status = 'processado' and (%(tipo)s = '' or tipo = %(tipo)s)
                   and coalesce(processado_em, recebido_em) < now() - interval '30 days'""", filtro)[0]["n"] if recentes else 0
    contagem = q("select status, count(*) as n from fin_app.documentos group by status")
    return {"documentos": docs, "contagem": {r["status"]: r["n"] for r in contagem}, "ocultos": ocultos}


@router.get("/contagem")
def contagem():
    return q("select count(*) filter (where status = 'novo') as novos from fin_app.documentos")[0]


def _doc(did: uuid.UUID):
    r = q("select * from fin_app.documentos where id = %s", (did,))
    if not r:
        raise HTTPException(404, "Documento não encontrado")
    return r[0]


def _arquivo(did) -> str:
    pasta = os.path.join(DOCS_DIR, str(uuid.UUID(str(did))))
    nome = next((f for f in os.listdir(pasta) if f.startswith("original.")), None) if os.path.isdir(pasta) else None
    if not nome:
        raise HTTPException(404, "Arquivo não encontrado")
    return os.path.join(pasta, nome)


@router.get("/{did}")
def detalhe(did: uuid.UUID):
    d = _doc(did)
    nfe = (d["dados"] or {}).get("nfe")
    if nfe:
        doc = re.sub(r"\D", "", nfe["emitente"]["cnpj"] or "")
        pessoa = q("""select id, nome from fin_app.pessoas where ativo and documento is not null
                      and regexp_replace(documento, '\\D', '', 'g') = %s limit 1""", (doc,)) if doc else []
        d["emitente_cadastrado"] = pessoa[0] if pessoa else None
    if d["tipo"] == "nota_fiscal":
        chave = (nfe or {}).get("chave") or (d["dados"] or {}).get("chave")
        d["compra"] = q("""select id, conciliada_em, finalidade from fin_app.compras where documento_id = %s or (%s::text is not null and chave = %s)""",
                        (did, chave, chave))
    d["no_financeiro"] = bool(d.get("lancamento_ids"))
    d["tem_previa"] = os.path.isfile(os.path.join(DOCS_DIR, str(did), "previa.png"))
    return d


@router.get("/{did}/arquivo")
def baixar(did: uuid.UUID):
    d = _doc(did)
    return FileResponse(_arquivo(did), media_type=d["mime"], filename=d["nome"], content_disposition_type="inline")


@router.get("/{did}/previa")
def previa(did: uuid.UUID):
    caminho = os.path.join(DOCS_DIR, str(did), "previa.png")
    if not os.path.isfile(caminho):
        raise HTTPException(404, "Sem prévia")
    return FileResponse(caminho, media_type="image/png")


class Status(BaseModel):
    status: str = Field(pattern="^(novo|processado|arquivado)$")


class Tipo(BaseModel):
    tipo: str = Field(pattern="^(boleto|nota_fiscal|relatorio_vendas|outro)$")


ROT_TIPO = {"boleto": "Boleto", "nota_fiscal": "Nota fiscal", "relatorio_vendas": "Relatório de vendas", "outro": "Outro"}


@router.post("/{did}/tipo")
def mudar_tipo(did: uuid.UUID, t: Tipo):
    """Documento que o sistema não reconheceu: o usuário diz o que é (o resumo genérico vira o nome do tipo)."""
    with psycopg.connect(**DSN) as conn:
        cur = conn.execute("""update fin_app.documentos set tipo = %s,
                                resumo = case when resumo is null or resumo = 'Documento não identificado' then %s else resumo end
                              where id = %s""", (t.tipo, f"{ROT_TIPO[t.tipo]} (marcado à mão)", did))
        if cur.rowcount != 1:
            raise HTTPException(404, "Documento não encontrado")
    return {"ok": True}


@router.post("/{did}/status")
def mudar_status(did: uuid.UUID, s: Status):
    with psycopg.connect(**DSN) as conn:
        cur = conn.execute("""update fin_app.documentos set status = %s,
                                processado_em = case when %s = 'processado' then now() else processado_em end
                              where id = %s""", (s.status, s.status, did))
        if cur.rowcount != 1:
            raise HTTPException(404, "Documento não encontrado")
    return {"ok": True}


@router.post("/{did}/leitor")
def abrir_no_leitor(did: uuid.UUID):
    """Abre boleto/NF em PDF ou foto no leitor (OCR com palavras clicáveis + linha digitável + IA)."""
    d = _doc(did)
    with open(_arquivo(did), "rb") as f:
        return boleto.processar(f.read(), d["nome"], documento_id=did)


class Parcela(BaseModel):
    vencimento: dt.date
    valor: Decimal = Field(gt=0, max_digits=14, decimal_places=2)


class LancarNfe(BaseModel):
    contato_id: uuid.UUID
    categoria_id: uuid.UUID | None = None
    conta_id: uuid.UUID | None = None
    descricao: str | None = Field(None, max_length=300)
    parcelas: list[Parcela] = Field(min_length=1, max_length=48)


@router.post("/{did}/lancar-nfe")
def lancar_nfe(did: uuid.UUID, b: LancarNfe):
    d = _doc(did)
    nfe = (d["dados"] or {}).get("nfe")
    if not nfe:
        raise HTTPException(422, "Este documento não é uma NF-e em XML")
    with psycopg.connect(**DSN, row_factory=dict_row) as conn, conn.transaction():
        cur = conn.cursor()
        if d.get("lancamento_ids"):
            raise HTTPException(409, "Esta nota já foi lançada no financeiro")
        cat_nome, conta_nome, contato = fin._nomes(cur, "pagar", b.categoria_id, b.conta_id, b.contato_id)
        # a conciliação de itens pode ter criado a compra antes: reaproveita (só liga o fornecedor)
        existente = cur.execute("""select id from fin_app.compras where documento_id = %s or (%s::text is not null and chave = %s)""",
                                (did, nfe.get("chave"), nfe.get("chave"))).fetchone()
        if existente:
            compra = existente["id"]
            cur.execute("update fin_app.compras set pessoa_id = %s where id = %s", (b.contato_id, compra))
        else:
            compra = cur.execute("""insert into fin_app.compras (documento_id, chave, numero, serie, emissao, emitente_cnpj,
                                        emitente_nome, pessoa_id, valor_produtos, valor_frete, valor_desconto, valor_total)
                                    values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) returning id""",
                                 (did, nfe.get("chave"), nfe["numero"], nfe["serie"], nfe["emissao"], nfe["emitente"]["cnpj"],
                                  nfe["emitente"]["nome"], b.contato_id, nfe["valor_produtos"], nfe["valor_frete"],
                                  nfe["valor_desconto"], nfe["valor_total"])).fetchone()["id"]
            cur.executemany("""insert into fin_app.compras_itens (compra_id, n, codigo, ean, descricao, ncm, cfop, unidade,
                                   quantidade, valor_unitario, valor_total)
                               values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                            [(compra, i["n"], i["codigo"], i["ean"], i["descricao"][:300], i["ncm"], i["cfop"], i["unidade"],
                              i["quantidade"], i["valor_unitario"], i["valor_total"]) for i in nfe["itens"]])
        n = len(b.parcelas)
        grupo = uuid.uuid4() if n > 1 else None
        desc = (b.descricao or f"NF {nfe['numero']} - {nfe['emitente']['fantasia'] or nfe['emitente']['nome']}").strip()
        ids = []
        for i, p in enumerate(b.parcelas, start=1):
            linha = {"tipo": "pagar", "vencimento": p.vencimento, "competencia": nfe["emissao"], "previsto_para": None,
                     "contato": contato, "contato_id": b.contato_id, "categoria_id": b.categoria_id,
                     "categoria_nome": cat_nome, "conta_id": b.conta_id, "conta_nome": conta_nome,
                     "descricao": f"{desc} ({i}/{n})" if n > 1 else desc, "valor": p.valor, "status": "aberto",
                     "valor_pago": None, "data_pagamento": None, "grupo_id": grupo, "parcela": i if grupo else None,
                     "total_parcelas": n if grupo else None, "recorrencia": "parcelado" if grupo else None}
            cur.execute(f"insert into fin_app.lancamentos ({', '.join(fin.COLS)}) "
                        f"values ({', '.join('%(' + c + ')s' for c in fin.COLS)}) returning id", linha)
            ids.append(cur.fetchone()["id"])
        cur.execute("""update fin_app.documentos set status = 'processado', processado_em = now(), lancamento_ids = %s
                       where id = %s""", (ids, did))
    return {"compra_id": str(compra), "lancamentos": [str(i) for i in ids]}
